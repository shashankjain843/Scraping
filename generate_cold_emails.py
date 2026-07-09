import os
import csv
import time
import json
import sys
import re
import requests
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
from email.header import Header
from docx import Document
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# Reconfigure stdout to use UTF-8 to prevent UnicodeEncodeError in Windows terminals
if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

# Configurations
MODEL_NAME = "openai/gpt-4o-mini"
OPENROUTER_API_URL = "https://openrouter.ai/api/v1/chat/completions"
TRACKER_FILE = "sent_emails_tracker.json"
MAX_DAILY_EMAILS = 10  # Enforces a limit of 10 emails per day

# Resume file path
RESUME_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "my_resume.docx")
RESUME_DATA = None

# Fallback resume
RESUME_DATA_FALLBACK = {
    "Name": "Shashank Jain",
    "Role": "Data Analyst",
    "Skills": [
        "Python", "SQL", "Pandas", "Numpy", "Data Cleaning", "Exploratory Data Analysis (EDA)", "Insight Generation",
        "Power BI", "Tableau", "PostgreSQL", "MongoDB", "FastAPI", "LangChain", "RAG", "Git", "JupyterNotebook", "VS Code", "SDLC",
        "Communication", "Problem Solving", "Teamwork", "Adaptability"
    ],
    "Experience": [
        {
            "Role": "Data Analyst Intern",
            "Company": "Appic Software Development (On-Site)",
            "Details": "Performed exploratory data analysis and preprocessing on 10,000+ rows using Pandas and NumPy; Applied data cleaning and transformation techniques improving dataset quality by 25%; Created interactive Power BI dashboards for KPI tracking reducing manual reporting effort by 30%."
        }
    ],
    "Projects": [
        {
            "Name": "Sales and Revenue Performance Analysis",
            "Technologies": ["Python", "Numpy", "SQL", "Streamlit", "Seaborn"],
            "Details": "Built a Python SQL ETL pipeline to clean 100k+ transactions and run YoY sales and cohort retention models in SQLite; Developed a Holt's linear forecasting model (31.6% MAPE) and RFM customer segmentation to analyze lifetime value; Developed an interactive Streamlit dashboard that identified a 30% discount tipping point, projecting a 1.8% gross profit margin increase."
        },
        {
            "Name": "Healthcare Analytics Dashboard",
            "Technologies": ["Python", "Pandas", "PostgreSQL", "XGBoost", "Streamlit"],
            "Details": "Designed a healthcare analytics system that predicts 30-day patient readmission risk using clinical and demographic patient data; Built a disease outbreak forecasting module analyzing historical case trends to predict future outbreak patterns by region and time; Developed an interactive dashboard comparing treatment effectiveness across patient groups using survival analysis to support hospital decision-making."
        }
    ],
    "Contact": {
        "Phone": "+91-7878927128",
        "LinkedIn": "linkedin.com/in/shashankjain",
        "GitHub": "github.com/shashankjain843"
    }
}

# Stats tracking for logs
stats = {
    "jobs_scraped": 0,
    "emails_generated": 0,
    "emails_sent": 0
}

def log_to_file(message):
    timestamp = time.strftime("%Y-%m-%d %H:%M:%S")
    formatted_msg = f"[{timestamp}] {message}\n"
    print(formatted_msg.strip())
    try:
        with open("logs.txt", "a", encoding="utf-8") as f:
            f.write(formatted_msg)
    except Exception as e:
        print(f"[ERROR] Failed to write to log file logs.txt: {e}")

def get_daily_sent_count():
    today = time.strftime("%Y-%m-%d")
    if os.path.exists(TRACKER_FILE):
        try:
            with open(TRACKER_FILE, "r") as f:
                data = json.load(f)
                if data.get("date") == today:
                    return data.get("count", 0)
        except Exception:
            pass
    return 0

def increment_daily_sent_count():
    today = time.strftime("%Y-%m-%d")
    count = get_daily_sent_count() + 1
    try:
        with open(TRACKER_FILE, "w") as f:
            json.dump({"date": today, "count": count}, f)
    except Exception as e:
        log_to_file(f"[WARNING] Failed to write to tracker file: {e}")
    return count

def find_email_in_text(text):
    if not text:
        return None
    emails = re.findall(r'\b[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}\b', text)
    for email in emails:
        if not any(x in email.lower() for x in ["noreply", "no-reply", "donotreply", "support@", "privacy@", "info@"]):
            return email
    return emails[0] if emails else None

def extract_subject_and_body(email_text):
    lines = email_text.strip().split("\n")
    subject = "Application for Data Analyst Position"
    body_lines = []
    subject_found = False
    
    for line in lines:
        if not subject_found and line.lower().startswith("subject:"):
            subject = line[len("subject:"):].strip()
            subject_found = True
        else:
            body_lines.append(line)
            
    body = "\n".join(body_lines).strip()
    return subject, body

def convert_docx_to_pdf(docx_path, pdf_path):
    try:
        from docx2pdf import convert
        convert(docx_path, pdf_path)
        log_to_file(f"[SUCCESS] Converted {docx_path} to {pdf_path} using docx2pdf.")
        return True
    except Exception as e:
        log_to_file(f"[WARNING] Failed to convert docx to pdf using docx2pdf: {e}. Trying raw win32com...")
        
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(docx_path)
        # Format 17 is for PDF
        doc.SaveAs(pdf_path, FileFormat=17)
        doc.Close()
        word.Quit()
        log_to_file(f"[SUCCESS] Converted {docx_path} to {pdf_path} using Word COM.")
        return True
    except Exception as e:
        log_to_file(f"[ERROR] Failed to convert docx to pdf using Word COM: {e}.")
        
    return False

def send_email_via_smtp(subject, body, recipient_email):
    global stats
    smtp_server = os.environ.get("SMTP_SERVER", "smtp.gmail.com")
    smtp_port = int(os.environ.get("SMTP_PORT", "587"))
    smtp_email = os.environ.get("SMTP_EMAIL")
    smtp_password = os.environ.get("SMTP_PASSWORD")
    
    if not smtp_email or not smtp_password:
        log_to_file("[ERROR] SMTP credentials not set in .env. Skipping email sending.")
        return False
        
    try:
        # Create a multipart message
        msg = MIMEMultipart()
        msg["Subject"] = Header(subject, "utf-8")
        msg["From"] = smtp_email
        msg["To"] = recipient_email
        
        # Attach body
        msg.attach(MIMEText(body, "plain", "utf-8"))
        
        dir_path = os.path.dirname(os.path.abspath(__file__))
        
        # Look for PDF resume first
        pdf_path = os.path.join(dir_path, "my_resume.pdf")
        docx_path = os.path.join(dir_path, "my_resume.docx")
        
        resume_path = None
        if os.path.exists(pdf_path):
            resume_path = pdf_path
        elif os.path.exists(docx_path):
            log_to_file("[INFO] PDF resume not found. Attempting to convert my_resume.docx to PDF...")
            success = convert_docx_to_pdf(docx_path, pdf_path)
            if success and os.path.exists(pdf_path):
                resume_path = pdf_path
            else:
                log_to_file("[WARNING] DOCX to PDF conversion failed. Falling back to original DOCX.")
                resume_path = docx_path
                
        if resume_path:
            filename = os.path.basename(resume_path)
            try:
                with open(resume_path, "rb") as attachment:
                    part = MIMEBase("application", "octet-stream")
                    part.set_payload(attachment.read())
                encoders.encode_base64(part)
                part.add_header(
                    "Content-Disposition",
                    f"attachment; filename= {filename}",
                )
                msg.attach(part)
                log_to_file(f"[INFO] Attached resume: {filename} to email.")
            except Exception as att_err:
                log_to_file(f"[WARNING] Failed to attach resume: {att_err}")
        else:
            log_to_file(f"[WARNING] Resume file (my_resume.pdf or my_resume.docx) not found in {dir_path}. Sending without attachment.")
            
        server = smtplib.SMTP(smtp_server, smtp_port, timeout=30)
        server.starttls()
        server.login(smtp_email, smtp_password)
        server.sendmail(smtp_email, [recipient_email], msg.as_string())
        server.quit()
        
        log_to_file(f"[SUCCESS] SMTP Email successfully sent to: {recipient_email}")
        stats["emails_sent"] += 1
        return True
    except Exception as smtp_err:
        log_to_file(f"[ERROR] SMTP sending failed to {recipient_email}: {smtp_err}")
        return False

def parse_resume_from_doc(filepath, api_key):
    try:
        doc = Document(filepath)
        raw_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    except Exception as docx_err:
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                raw_text = f.read().strip()
        except Exception as txt_err:
            log_to_file(f"[ERROR] Resume file read failed (Docx error: {docx_err}, Plain text error: {txt_err})")
            return None

    if not raw_text.strip():
        log_to_file("[ERROR] Resume file empty or readable text not found!")
        return None

    log_to_file(f"[INFO] Resume read successfully from: {filepath}")
    log_to_file("[INFO] Parsing resume using AI parser...")

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }

    prompt = f"""Neeche resume ka raw text hai. Isse ek structured JSON mein convert karo.
JSON ka format bilkul yeh hona chahiye (koi extra text ya explanation nahi, sirf JSON):
{{
    "Name": "...",
    "Role": "...",
    "Skills": ["skill1", "skill2"],
    "Experience": [
        {{"Role": "...", "Company": "...", "Details": "..."}}
    ],
    "Projects": [
        {{"Name": "...", "Technologies": ["tech1", "tech2"], "Details": "..."}}
    ],
    "Contact": {{"Phone": "...", "LinkedIn": "...", "GitHub": "..."}}
}}

Resume Text:
{raw_text}
"""

    payload = {
        "model": MODEL_NAME,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.1
    }

    try:
        response = requests.post(OPENROUTER_API_URL, headers=headers, json=payload, timeout=30)
        response.raise_for_status()
        content = response.json()["choices"][0]["message"]["content"].strip()

        if content.startswith("```"):
            lines = content.split("\n")
            lines = [l for l in lines if not l.strip().startswith("```")]
            content = "\n".join(lines).strip()

        resume_data = json.loads(content)
        log_to_file("[INFO] Resume parsed successfully by AI.")
        return resume_data
    except Exception as e:
        log_to_file(f"[ERROR] Failed parsing resume with AI: {e}")
        return None

def get_api_key():
    api_key = os.environ.get("OPENROUTER_API_KEY")
    if not api_key:
        log_to_file("[ERROR] OPENROUTER_API_KEY environment variable is not set in environment or .env!")
        api_key = input("Please paste your OpenRouter API Key: ").strip()
        if not api_key:
            sys.exit(1)
    return api_key

def generate_cold_email(api_key, company_name, job_title, job_description):
    global stats
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "https://github.com/shashankjain843",
        "X-Title": "Cold Email Generator"
    }

    system_prompt = (
        "Tu ek professional cold email writer hai. Email short, direct, no fluff, 5-6 lines ka ho. "
        "Candidate ke resume se sirf wahi skills/projects mention karo jo job description se match karte hain. "
        "Generic buzzwords mat use karo. Write the email in English. "
        "Do not use placeholders like [Company Name], [Job Title], [Your Name], or any brackets. "
        "Always replace them with the actual names provided. "
        "Always address the email to 'Dear Hiring Team,' or 'Dear [Company Name] Hiring Team,' instead of 'Dear Hiring Manager,' or 'Dear Recruiter,'. "
        "The response must start directly with 'Subject: [Subject Line]' followed by the email body. "
        "Do not write any introductory or concluding conversation text, just start with Subject:.\n"
        "CRITICAL: Do NOT use any spam-trigger words or phrases such as 'Free', 'Guaranteed', "
        "'Act now', 'Limited time', 'Click here', 'Risk-free', 'Special offer', 'Click below', or 'Hurry'. "
        "Subject lines should be unique, professional, and personalized to the role and company."
    )

    user_prompt = f"""
Candidate Resume:
- Name: {RESUME_DATA['Name']}
- Target Role: {RESUME_DATA['Role']}
- Contact: Phone: {RESUME_DATA['Contact']['Phone']}, LinkedIn: {RESUME_DATA['Contact']['LinkedIn']}, GitHub: {RESUME_DATA['Contact']['GitHub']}
- Skills: {', '.join(RESUME_DATA['Skills'])}
- Experience: Data Analyst Intern at {RESUME_DATA['Experience'][0]['Company']} ({RESUME_DATA['Experience'][0]['Details']})
- Projects:
  1. {RESUME_DATA['Projects'][0]['Name']} (Technologies: {', '.join(RESUME_DATA['Projects'][0]['Technologies'])})
  2. {RESUME_DATA['Projects'][1]['Name']} (Technologies: {', '.join(RESUME_DATA['Projects'][1]['Technologies'])})

Job Listing details:
- Company Name: {company_name}
- Job Title: {job_title}
- Job Description/Required Skills: {job_description}

Write a personalized cold email from {RESUME_DATA['Name']} to the hiring team at {company_name} for the '{job_title}' position. Match relevant skills and projects from his resume. Make it short (5-6 lines), professional, and direct. Address the email with 'Dear Hiring Team,'. Start with the Subject line.
"""

    payload = {
        "model": MODEL_NAME,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        "temperature": 0.7
    }

    try:
        response = requests.post(OPENROUTER_API_URL, headers=headers, json=payload, timeout=30)
        response.raise_for_status()
        data = response.json()
        if "choices" in data and len(data["choices"]) > 0:
            stats["emails_generated"] += 1
            email_text = data["choices"][0]["message"]["content"].strip()
            # Post-process to ensure "Dear Hiring Team" greeting
            email_text = re.sub(r'Dear\s+Hiring\s+Manager\b', 'Dear Hiring Team', email_text, flags=re.IGNORECASE)
            email_text = re.sub(r'Dear\s+Recruiter\b', 'Dear Hiring Team', email_text, flags=re.IGNORECASE)
            return email_text
        else:
            log_to_file(f"[ERROR] Invalid API response format from OpenRouter: {data}")
            return None
    except Exception as e:
        log_to_file(f"[ERROR] AI cold email generation API call failed: {e}")
        return None

def run_mode_a(api_key):
    log_to_file("\n--- MODE A: Standalone (Manual Cold Email Generator) ---")
    company_name = input("Enter Company Name: ").strip()
    job_title = input("Enter Job Title: ").strip()

    print("\nEnter Job Description / Required Skills.")
    print("Options:")
    print("1. Paste/type text directly (Press Enter, then type 'DONE' on a new line and press Enter to finish)")
    print("2. Type 'FILE:<filepath>' to read the description from a text file (e.g. FILE:desc.txt)")
    
    lines = []
    first_input = input("Input choice or first line: ").strip()
    
    if first_input.upper().startswith("FILE:"):
        file_path = first_input[5:].strip()
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                job_desc = f.read().strip()
            log_to_file(f"Successfully read job description from: {file_path}")
        except Exception as e:
            log_to_file(f"[ERROR] Could not read file: {e}. Exiting manual entry.")
            return
    else:
        if first_input.strip().upper() != "DONE":
            lines.append(first_input)
        while True:
            line = input()
            if line.strip().upper() == "DONE":
                break
            lines.append(line)
        job_desc = "\n".join(lines)

    if not company_name or not job_title or not job_desc:
        log_to_file("[ERROR] Company Name, Job Title, and Job Description cannot be empty!")
        return

    log_to_file("Generating email template using OpenRouter API...")
    email_text = generate_cold_email(api_key, company_name, job_title, job_desc)
    
    if email_text:
        print("\n" + "="*50)
        print("GENERATED COLD EMAIL:")
        print("="*50)
        print(email_text)
        print("="*50)
        
        # Save to text file
        out_filename = "cold_email_output.txt"
        try:
            with open(out_filename, "w", encoding="utf-8") as f:
                f.write(f"Company: {company_name}\nJob Title: {job_title}\n\n{email_text}\n")
            log_to_file(f"[SUCCESS] Email successfully saved to: {out_filename}")
        except Exception as e:
            log_to_file(f"[WARNING] Failed to save draft to file: {e}")
            
        recipient_email = input("\nEnter recipient's email address to send now (leave blank to skip sending): ").strip()
        if recipient_email:
            daily_sent = get_daily_sent_count()
            if daily_sent >= MAX_DAILY_EMAILS:
                log_to_file(f"[LIMIT REACHED] Maximum daily sending limit of {MAX_DAILY_EMAILS} emails has been reached. Please run the script again tomorrow to continue sending.")
                return
            
            subject, body = extract_subject_and_body(email_text)
            sent_ok = send_email_via_smtp(subject, body, recipient_email)
            if sent_ok:
                increment_daily_sent_count()
    else:
        log_to_file("[ERROR] Could not generate cold email.")

def run_mode_b(api_key):
    log_to_file("\n--- MODE B: Scraper Integration (Automatic Cold Email Generator) ---")
    
    files = [f for f in os.listdir(".") if f.endswith(".csv")]
    print("\nAvailable CSV files in current directory:")
    for idx, f in enumerate(files, 1):
        print(f" {idx}. {f}")
    
    default_file = "linkedin_fresher_data_analyst_jobs_merged.csv"
    csv_path = input(f"\nEnter the path to the jobs CSV file (default: {default_file}): ").strip()
    if not csv_path:
        csv_path = default_file

    if not os.path.exists(csv_path):
        log_to_file(f"[ERROR] File '{csv_path}' does not exist!")
        return

    if "naukri" in os.path.basename(csv_path).lower():
        log_to_file(f"[INFO] Skipping file '{csv_path}' because it belongs to Naukri.com.")
        return

    log_to_file(f"Reading jobs from: {csv_path}...")
    jobs = []
    try:
        with open(csv_path, mode="r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames
            if not headers:
                log_to_file("[ERROR] CSV file is empty or has no headers!")
                return
                
            def find_col(possible_names):
                for name in possible_names:
                    for header in headers:
                        if name.lower() == header.lower().strip():
                            return header
                return None
            
            company_col = find_col(["Company Name", "Company", "company_name", "company"])
            title_col = find_col(["Job Title", "Title", "job_title", "title"])
            skills_col = find_col(["Required Skills", "Job Description", "skills", "description", "required_skills"])
            url_col = find_col(["Job Posting Link / URL", "Job URL", "url", "link", "job_url", "posting_url"])
            source_col = find_col(["Source Platform", "source", "Source"])
            email_col = find_col(["Recruiter Email", "Email", "recruiter_email", "email", "contact_email"])

            if not company_col or not title_col:
                log_to_file(f"[ERROR] Could not map key columns. Found columns: {headers}")
                return

            for row in reader:
                if source_col:
                    source_val = (row.get(source_col) or "").strip().lower()
                    if "naukri" in source_val:
                        continue
                
                comp = (row.get(company_col) or "").strip()
                title = (row.get(title_col) or "").strip()
                skills = (row.get(skills_col) or "").strip()
                url = (row.get(url_col) or "Not Available").strip()
                email = (row.get(email_col) or "").strip() if email_col else ""
                
                if comp and title:
                    jobs.append({
                        "company": comp,
                        "title": title,
                        "skills": skills,
                        "url": url,
                        "email": email
                    })
    except Exception as e:
        log_to_file(f"[ERROR] Failed to read CSV file: {e}")
        return

    if not jobs:
        log_to_file("[INFO] No LinkedIn jobs found to process in the CSV.")
        return

    log_to_file(f"Found {len(jobs)} LinkedIn job(s) to process.")
    output_file = "generated_cold_emails.csv"
    
    mode = "w"
    write_header = True
    if os.path.exists(output_file):
        ans = input(f"Output file '{output_file}' already exists. (A)ppend, (O)verwrite, or (C)ancel? ").strip().upper()
        if ans == "A":
            mode = "a"
            write_header = False
        elif ans == "C":
            log_to_file("Operation cancelled.")
            return

    processed_count = 0
    success_count = 0

    try:
        with open(output_file, mode=mode, encoding="utf-8-sig", newline="") as out_f:
            writer = csv.writer(out_f)
            if write_header:
                writer.writerow(["Company", "Job Title", "Generated Email", "Job URL", "Recipient Email"])

            for idx, job in enumerate(jobs, 1):
                log_to_file(f"\n[{idx}/{len(jobs)}] Processing: {job['title']} at {job['company']}")
                
                email_text = generate_cold_email(api_key, job['company'], job['title'], job['skills'])
                processed_count += 1
                
                if email_text:
                    recipient = job['email'] if job['email'] else find_email_in_text(job['skills'])
                    
                    writer.writerow([job['company'], job['title'], email_text, job['url'], recipient or "Not Found"])
                    out_f.flush()
                    success_count += 1
                    log_to_file(f"  [SUCCESS] Cold email generated and saved.")
                    
                    if recipient:
                        daily_sent = get_daily_sent_count()
                        if daily_sent >= MAX_DAILY_EMAILS:
                            log_to_file(f"\n[LIMIT REACHED] Daily sending limit of {MAX_DAILY_EMAILS} emails reached. Exiting sending loop. Drafts saved.")
                            break
                        
                        subject, body = extract_subject_and_body(email_text)
                        log_to_file(f"  [SMTP] Recruiter email found: {recipient}. Attempting SMTP sending...")
                        
                        sent_ok = send_email_via_smtp(subject, body, recipient)
                        if sent_ok:
                            new_count = increment_daily_sent_count()
                            log_to_file(f"  [LIMIT CHECK] Daily emails sent: {new_count}/{MAX_DAILY_EMAILS}")
                            
                            # Random delay 30-60 seconds for deliverability
                            if idx < len(jobs) and new_count < MAX_DAILY_EMAILS:
                                delay = random.randint(30, 60)
                                log_to_file(f"  [DELAY] Waiting {delay} seconds between email sends to protect sender reputation...")
                                time.sleep(delay)
                        else:
                            log_to_file("  [SMTP ERROR] Skipping delay since sending failed.")
                    else:
                        log_to_file("  [INFO] No recipient email found in job details. Draft saved only.")
                        
                        # Short delay between API requests to respect OpenRouter rate limits
                        if idx < len(jobs):
                            time.sleep(2.5)
                else:
                    log_to_file(f"  [SKIPPED] Failed to generate email for this job.")

        log_to_file(f"\n[FINISHED] Processed {processed_count} jobs. Successfully generated {success_count} emails. Total emails sent via SMTP: {stats['emails_sent']}")
        log_to_file(f"Results saved to: {output_file}")
    except Exception as e:
        log_to_file(f"[ERROR] Failed writing to output file: {e}")

def main():
    global RESUME_DATA

    log_to_file("="*60)
    log_to_file("AUTOMATIC COLD EMAIL GENERATOR (OPENROUTER)")
    log_to_file("="*60)

    api_key = get_api_key()

    if os.path.exists(RESUME_FILE):
        RESUME_DATA = parse_resume_from_doc(RESUME_FILE, api_key)
        if RESUME_DATA is None:
            log_to_file("[WARNING] Resume parse failed. Using fallback (hardcoded) resume.")
            RESUME_DATA = RESUME_DATA_FALLBACK
    else:
        log_to_file(f"[WARNING] '{RESUME_FILE}' not found! Using fallback (hardcoded) resume.")
        RESUME_DATA = RESUME_DATA_FALLBACK

    log_to_file(f"[READY] Resume loaded for: {RESUME_DATA.get('Name', 'Unknown')} - {RESUME_DATA.get('Role', 'Unknown')}")

    print("\nSelect Mode:")
    print("1. Mode A – Standalone (Enter job details manually)")
    print("2. Mode B – Scraper Integration (Process LinkedIn jobs from CSV)")

    choice = input("Enter choice (1 or 2): ").strip()
    if choice == "1":
        run_mode_a(api_key)
    elif choice == "2":
        run_mode_b(api_key)
    else:
        log_to_file("Invalid choice. Exiting.")

if __name__ == "__main__":
    main()
